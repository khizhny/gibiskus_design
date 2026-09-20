from __future__ import annotations

import argparse
import os
import sqlite3
import sys
import tempfile
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it or run with the bundled Codex Python runtime."
    ) from exc


TYPE_ROWS = [
    (1, "\u0410\u0432\u0442\u043e\u043d\u043e\u043c\u043d\u0430 \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0456\u043a\u0430 \u041a\u0440\u0438\u043c, \u043e\u0431\u043b\u0430\u0441\u0442\u0456", "O"),
    (2, "\u043c\u0456\u0441\u0442\u0430, \u0449\u043e \u043c\u0430\u044e\u0442\u044c \u0441\u043f\u0435\u0446\u0456\u0430\u043b\u044c\u043d\u0438\u0439 \u0441\u0442\u0430\u0442\u0443\u0441", "K"),
    (3, "\u0440\u0430\u0439\u043e\u043d\u0438 \u0432 \u043e\u0431\u043b\u0430\u0441\u0442\u044f\u0445 \u0442\u0430 \u0410\u0432\u0442\u043e\u043d\u043e\u043c\u043d\u0456\u0439 \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0456\u0446\u0456 \u041a\u0440\u0438\u043c", "P"),
    (4, "\u0442\u0435\u0440\u0438\u0442\u043e\u0440\u0456\u0457 \u0442\u0435\u0440\u0438\u0442\u043e\u0440\u0456\u0430\u043b\u044c\u043d\u0438\u0445 \u0433\u0440\u043e\u043c\u0430\u0434 (\u043d\u0430\u0437\u0432\u0438 \u0442\u0435\u0440\u0438\u0442\u043e\u0440\u0456\u0430\u043b\u044c\u043d\u0438\u0445 \u0433\u0440\u043e\u043c\u0430\u0434) \u0432 \u043e\u0431\u043b\u0430\u0441\u0442\u044f\u0445, \u0442\u0435\u0440\u0438\u0442\u043e\u0440\u0456\u0430\u043b\u044c\u043d\u0456 \u0433\u0440\u043e\u043c\u0430\u0434\u0438 \u0410\u0432\u0442\u043e\u043d\u043e\u043c\u043d\u043e\u0457 \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0456\u043a\u0438 \u041a\u0440\u0438\u043c", "H"),
    (5, "\u043c\u0456\u0441\u0442\u0430", "M"),
    (6, "\u0441\u0435\u043b\u0438\u0449\u0430 \u043c\u0456\u0441\u044c\u043a\u043e\u0433\u043e \u0442\u0438\u043f\u0443", "T"),
    (7, "\u0441\u0435\u043b\u0430", "C"),
    (8, "\u0441\u0435\u043b\u0438\u0449\u0430", "X"),
    (9, "\u0440\u0430\u0439\u043e\u043d\u0438 \u0432 \u043c\u0456\u0441\u0442\u0430\u0445", "B"),
]

SOURCE_TO_TYPE_ID = {code: type_id for type_id, _, code in TYPE_ROWS}
LEVEL_COLS = (0, 1, 2, 3, 4)
DEFAULT_HEADER_ROW = 3
ROOT_TYPE = "_T"
LEVEL_BY_TYPE = {
    "O": 1,
    "K": 1,
    "P": 2,
    "H": 3,
    "M": 4,
    "T": 4,
    "C": 4,
    "X": 4,
    "B": 5,
}
FLAT_HEADERS = {"Structure URN", "ID", "Name UK", "KATOTTG_Category", "Parent ID"}


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    database_dir = script_dir.parent / "database"
    parser = argparse.ArgumentParser(
        description="Convert an official KATOTTG XLSX export to map.sqlite."
    )
    parser.add_argument(
        "xlsx",
        nargs="?",
        type=Path,
        default=database_dir / "map.xlsx",
        help="Input XLSX path. Defaults to database/map.xlsx.",
    )
    parser.add_argument(
        "sqlite",
        nargs="?",
        type=Path,
        default=database_dir / "map.sqlite",
        help="Output SQLite path. Defaults to database/map.sqlite.",
    )
    return parser.parse_args()


def find_data_start_row(ws) -> int:
    category_header = "\u041a\u0430\u0442\u0435\u0433\u043e\u0440\u0456\u044f \u043e\u0431\u2019\u0454\u043a\u0442\u0430"
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if len(row) > 5 and row[5] is not None and str(row[5]).strip() == category_header:
            return row_index + 1
    return DEFAULT_HEADER_ROW + 1


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFC", str(value).strip())
    return text or None


def find_flat_header_row(ws) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(
        ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1
    ):
        headers = [clean(value) for value in row]
        if FLAT_HEADERS.issubset(headers):
            return row_index, {name: headers.index(name) for name in FLAT_HEADERS}
    return None


def build_chain(records: dict[str, dict[str, str | None]], entry_id: str):
    chain = []
    seen = set()
    current_id = entry_id
    while current_id is not None:
        if current_id in seen:
            raise SystemExit(f"Parent cycle detected at {current_id}")
        seen.add(current_id)
        item = records.get(current_id)
        if item is None:
            raise SystemExit(f"Missing source record for parent ID: {current_id}")
        chain.append(item)
        current_id = item["parent_id"]
    return list(reversed(chain))


def load_flat_entries(ws, header_row: int, columns: dict[str, int]):
    records: dict[str, dict[str, str | None]] = {}
    source_urns = set()
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        source_urn = clean(row[columns["Structure URN"]])
        if source_urn:
            source_urns.add(source_urn)
        entry_id = clean(row[columns["ID"]])
        category = clean(row[columns["KATOTTG_Category"]])
        if not entry_id or not category:
            continue
        if entry_id in records:
            raise SystemExit(f"Duplicate entry id: {entry_id}")
        if category != ROOT_TYPE and category not in SOURCE_TO_TYPE_ID:
            raise SystemExit(f"Unsupported KATOTTG category {category!r}: {entry_id}")
        records[entry_id] = {
            "id": entry_id,
            "type": category,
            "name": clean(row[columns["Name UK"]]) or "",
            "parent_id": clean(row[columns["Parent ID"]]),
        }

    if not records:
        raise SystemExit("No KATOTTG records found in the flat XLSX export")
    roots = [item for item in records.values() if item["type"] == ROOT_TYPE]
    if len(roots) != 1:
        raise SystemExit(f"Expected one {ROOT_TYPE} root record, found {len(roots)}")
    if len(source_urns) != 1:
        raise SystemExit(f"Expected one Structure URN, found {sorted(source_urns)}")
    for item in records.values():
        parent_id = item["parent_id"]
        if parent_id is not None and parent_id not in records:
            raise SystemExit(f"Missing parent {parent_id} for {item['id']}")

    entries = []
    for entry_id, item in records.items():
        category = item["type"]
        if category == ROOT_TYPE:
            continue
        levels: list[str | None] = [None, None, None, None, None]
        previous_id = None
        previous_level = 0
        for ancestor in build_chain(records, entry_id):
            ancestor_type = ancestor["type"]
            if ancestor_type == ROOT_TYPE:
                continue
            target_level = LEVEL_BY_TYPE[ancestor_type]
            if target_level <= previous_level:
                raise SystemExit(
                    f"Invalid hierarchy for {entry_id}: "
                    f"{ancestor_type} at level {target_level}"
                )
            if previous_id is not None:
                for level in range(previous_level + 1, target_level):
                    levels[level - 1] = previous_id
            levels[target_level - 1] = ancestor["id"]
            previous_id = ancestor["id"]
            previous_level = target_level

        populated = [value for value in levels if value]
        if not populated or populated[-1] != entry_id:
            raise SystemExit(f"Failed to encode hierarchy for {entry_id}: {levels}")
        entries.append((*levels, SOURCE_TO_TYPE_ID[category], item["name"]))
    return entries


def load_legacy_entries(ws):
    data_start_row = find_data_start_row(ws)
    entries = []
    seen_ids = set()
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        source_code = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
        type_id = SOURCE_TO_TYPE_ID.get(source_code)
        if type_id is None:
            continue

        levels = [
            str(row[i]).strip() if len(row) > i and row[i] is not None else None
            for i in LEVEL_COLS
        ]
        non_empty_levels = [(i, value) for i, value in enumerate(levels) if value]
        if not non_empty_levels:
            continue

        _, entry_id = non_empty_levels[-1]
        name = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""
        if entry_id in seen_ids:
            raise SystemExit(f"Duplicate entry id: {entry_id}")
        seen_ids.add(entry_id)
        entries.append((*levels, type_id, name))
    return entries


def load_entries(
    xlsx_path: Path,
) -> list[tuple[str | None, str | None, str | None, str | None, str | None, int, str]]:
    if not xlsx_path.exists():
        raise SystemExit(f"Input file not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        flat_header = find_flat_header_row(ws)
        if flat_header is not None:
            entries = load_flat_entries(ws, *flat_header)
        else:
            entries = load_legacy_entries(ws)
    finally:
        wb.close()

    if not entries:
        raise SystemExit(f"No supported records found in {xlsx_path}")

    return entries


def write_database(
    sqlite_path: Path,
    entries: list[tuple[str | None, str | None, str | None, str | None, str | None, int, str]],
) -> None:
    sqlite_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{sqlite_path.stem}.",
        suffix=".tmp.sqlite",
        dir=sqlite_path.parent,
    )
    os.close(fd)
    temp_path = Path(temp_name)

    try:
        conn = sqlite3.connect(temp_path)
        try:
            conn.execute("PRAGMA foreign_keys = ON")
            conn.executescript(
                """
                CREATE TABLE entry_types (
                    id INTEGER PRIMARY KEY,
                    name VARCHAR(150) NOT NULL CHECK (length(name) <= 150),
                    code VARCHAR(1) NOT NULL UNIQUE CHECK (length(code) = 1)
                );

                CREATE TABLE entries (
                    l1_parent_id TEXT,
                    l2_parent_id TEXT,
                    l3_parent_id TEXT,
                    l4_parent_id TEXT,
                    l5_parent_id TEXT,
                    type INTEGER NOT NULL,
                    name TEXT NOT NULL,
                    FOREIGN KEY (type) REFERENCES entry_types(id)
                );

                CREATE INDEX idx_entries_l1_parent_id ON entries(l1_parent_id);
                CREATE INDEX idx_entries_l2_parent_id ON entries(l2_parent_id);
                CREATE INDEX idx_entries_l3_parent_id ON entries(l3_parent_id);
                CREATE INDEX idx_entries_l4_parent_id ON entries(l4_parent_id);
                CREATE INDEX idx_entries_l5_parent_id ON entries(l5_parent_id);
                CREATE INDEX idx_entries_type ON entries(type);
                CREATE INDEX idx_entries_name ON entries(name);
                CREATE INDEX idx_entry_types_id ON entry_types(id);

                CREATE VIEW oblasti AS
                    SELECT l1_parent_id AS oblast_id, name FROM entries where l2_parent_id is null;
                """
            )
            conn.executemany(
                "INSERT INTO entry_types (id, name, code) VALUES (?, ?, ?)",
                TYPE_ROWS,
            )
            conn.executemany(
                """
                INSERT INTO entries (
                    l1_parent_id,
                    l2_parent_id,
                    l3_parent_id,
                    l4_parent_id,
                    l5_parent_id,
                    type,
                    name
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                entries,
            )
            fk_errors = conn.execute("PRAGMA foreign_key_check").fetchall()
            if fk_errors:
                raise SystemExit(f"Foreign key check failed: {fk_errors[:10]}")
            conn.commit()
        finally:
            conn.close()

        try:
            os.replace(temp_path, sqlite_path)
        except PermissionError:
            data = temp_path.read_bytes()
            mode = "r+b" if sqlite_path.exists() else "wb"
            with sqlite_path.open(mode) as target:
                target.write(data)
                target.truncate()
            temp_path.unlink()
    finally:
        if temp_path.exists():
            temp_path.unlink()


def print_summary(sqlite_path: Path) -> None:
    conn = sqlite3.connect(sqlite_path)
    try:
        rows = conn.execute(
            """
            SELECT et.code, COUNT(e.name)
            FROM entry_types et
            LEFT JOIN entries e ON e.type = et.id
            GROUP BY et.id
            ORDER BY et.id
            """
        ).fetchall()
        entry_count = conn.execute("SELECT COUNT(*) FROM entries").fetchone()[0]
        oblast_count = conn.execute("SELECT COUNT(*) FROM oblasti").fetchone()[0]
        print(f"Created {sqlite_path}")
        print(f"entries={entry_count}")
        print(f"oblasti={oblast_count}")
        print("types=" + ", ".join(f"{code}:{count}" for code, count in rows))
    finally:
        conn.close()


def main() -> int:
    args = parse_args()
    entries = load_entries(args.xlsx)
    write_database(args.sqlite, entries)
    print_summary(args.sqlite)
    return 0


if __name__ == "__main__":
    sys.exit(main())
